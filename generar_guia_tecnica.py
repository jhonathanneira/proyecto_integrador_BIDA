from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('output') / 'Guia_preguntas_tecnicas_BIDA.docx'
BLUE = '17365D'
GRAY = 'F2F2F2'
BORDER = 'D9D9D9'

def font(run, size=10.2, bold=False, color=None):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)

def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.first_child_found_in('w:tcBorders')
    if node is None:
        node = OxmlElement('w:tcBorders')
        tc_pr.append(node)
    for edge in ('top', 'left', 'bottom', 'right'):
        element = OxmlElement('w:' + edge)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), BORDER)
        node.append(element)

def margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement('w:tcMar')
    for side in ('top', 'start', 'bottom', 'end'):
        element = OxmlElement('w:' + side)
        element.set(qn('w:w'), '100')
        element.set(qn('w:type'), 'dxa')
        node.append(element)
    tc_pr.append(node)

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(13 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    font(r, 14 if level == 1 else 11.5, True)
    return p

def text(doc, content, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(content)
    font(r)
    return p

def qa(doc, question, answer, clue=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run('Pregunta: ' + question)
    font(r, bold=True, color=BLUE)
    a = doc.add_paragraph()
    a.paragraph_format.space_after = Pt(2 if clue else 6)
    a.paragraph_format.line_spacing = 1.08
    r = a.add_run('Respuesta: ' + answer)
    font(r)
    if clue:
        c = doc.add_paragraph()
        c.paragraph_format.left_indent = Inches(0.22)
        c.paragraph_format.first_line_indent = Inches(-0.22)
        c.paragraph_format.space_after = Pt(6)
        r = c.add_run('Codigo a mencionar: ' + clue)
        font(r, 9.2, False, '555555')

def add_table(doc, headers, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = 'Table Grid'
    for i, item in enumerate(headers):
        cell = tbl.rows[0].cells[i]
        shade(cell, BLUE); borders(cell); margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.width = Inches(widths[i])
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(item); font(r, 9.2, True, 'FFFFFF')
    tr_pr = tbl.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'), 'true'); tr_pr.append(repeat)
    for row_no, row in enumerate(rows):
        cells = tbl.add_row().cells
        for i, item in enumerate(row):
            cell = cells[i]
            if row_no % 2: shade(cell, GRAY)
            borders(cell); margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.width = Inches(widths[i])
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(item); font(r, 9.1)
            if i == 0: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return tbl

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.6)
sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)
for name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Aptos'
    style._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    style.font.color.rgb = RGBColor(0, 0, 0)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('BIDA  Guia tecnica de sustentacion')
font(r, 8, False, '666666')

p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run('Guia tecnica de preguntas sobre el codigo BIDA')
font(r, 21, True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(17)
r = p.add_run('Respuestas para explicar arquitectura decisiones de implementacion y mejoras')
font(r, 10.8, False, '444444')

text(doc, 'Esta guia no pretende memorizar codigo. Sirve para explicar con criterio que hace cada parte, por que se eligio esa solucion y que mejorarias si el proyecto avanzara a produccion. Cuando te pregunten, empieza por la idea y luego nombra el archivo o componente que la implementa.', 10)

heading(doc, 'Mapa rapido del proyecto', 1)
add_table(doc, ['Componente', 'Responsabilidad', 'Ejemplo para explicar'], [
    ['app.js', 'Configura Express, middleware, conexion a TiDB, servidor y rutas.', 'Es el punto de entrada del backend.'],
    ['routes', 'Define endpoints y reglas de cada modulo.', 'clienteRoutes consulta y administra pacientes.'],
    ['autorizacion y sesiones', 'Valida token y permisos por especialidad.', 'Administracion puede crear y editar recursos protegidos.'],
    ['public', 'Contiene las vistas y JavaScript del navegador.', 'administracion.js consume la API y muestra la agenda.'],
    ['database', 'Scripts de carga de datos para demostracion.', 'Las cargas son transaccionales e idempotentes.'],
    ['test', 'Pruebas automatizadas con Jest y Supertest.', 'Verifican CRUD, autorizacion y validaciones.']
], [.95, 2.25, 3.6])

heading(doc, 'Servidor y arquitectura', 1)
qa(doc, 'Por que app.js es importante?', 'Porque concentra la configuracion transversal: carga variables de entorno, crea Express, habilita CORS, publica archivos estaticos, crea el pool de TiDB y monta las rutas bajo /app. Asi cada archivo de ruta se enfoca en su modulo.', 'app.js, require dotenv, express.static y app.use /app')
qa(doc, 'Por que las rutas reciben connection como parametro?', 'Es una forma sencilla de inyeccion de dependencias. app.js crea una sola fuente de acceso a datos y cada modulo recibe esa dependencia. Esto facilita las pruebas porque se puede reemplazar por un mock.', 'productoRoutes(connection), agendaRoutes(connection) y mocks de mysql2 en test')
qa(doc, 'Que significa que la API sea REST?', 'Que los recursos se exponen mediante endpoints y verbos HTTP: GET consulta, POST crea, PUT actualiza y DELETE elimina. Por ejemplo, /app/cliente y /app/cliente/:id representan el recurso cliente.', 'clienteRoutes.js y productoRoutes.js')
qa(doc, 'Por que se usa express.static?', 'Para que el mismo servidor entregue los archivos del frontend desde public. Esto evita depender de un servidor distinto para HTML, CSS y JavaScript durante el despliegue.', 'app.use express.static(path.join(__dirname, public))')
qa(doc, 'Por que el servidor escucha en 0.0.0.0 y usa process.env.PORT?', 'Render entrega el puerto mediante una variable de entorno y necesita que la aplicacion escuche en todas las interfaces. Localmente el proyecto usa 3000 si PORT no esta definida.', 'app.listen(puerto, 0.0.0.0) y process.env.PORT')

heading(doc, 'Base de datos y consultas', 1)
qa(doc, 'Por que se cambio createConnection por createPool?', 'Con una conexion unica, una desconexion temporal de TiDB podia emitir PROTOCOL_CONNECTION_LOST y cerrar el proceso. El pool administra varias conexiones, las reutiliza y puede abrir una nueva cuando una deja de estar disponible.', 'mysql.createPool con waitForConnections, connectionLimit y enableKeepAlive')
qa(doc, 'El pool cambia las rutas existentes?', 'No de forma relevante. Tanto la conexion como el pool ofrecen query, por eso se conserva la variable connection y las rutas siguen llamando connection.query. El cambio fue localizado en app.js.', 'connection.query en los controladores')
qa(doc, 'Como previenen inyeccion SQL?', 'Las consultas que incluyen datos variables usan placeholders ? y un arreglo de parametros. mysql2 se encarga de enviar los valores separados de la sentencia SQL.', 'WHERE idCliente = ? y arrays como req.params.id')
qa(doc, 'Por que algunas respuestas tienen total y lista?', 'Es una convención del API para entregar el conjunto de resultados y su cantidad. Por ejemplo, cliente devuelve total y clientes; tratamiento devuelve total y tratamientos. El frontend toma la propiedad lista correspondiente.', 'clienteRoutes.js y tratamientoRoutes.js')
qa(doc, 'Como se relaciona una cita con el medico?', 'La columna idEmpleado se guarda en cita. Al consultar, la agenda hace un LEFT JOIN con empleado para mostrar nombreEmpleado sin perder una cita aunque el nombre no estuviera disponible.', 'agendaRoutes.js, SELECT C y LEFT JOIN empleado E')
qa(doc, 'Por que las cargas de datos usan transacciones?', 'Porque una carga debe ser atomica: si falla un producto, una cita o una relacion, se hace rollback y no queda una base parcialmente cargada. Si todo funciona, se hace commit.', 'beginTransaction, commit y rollback en los scripts de database')
qa(doc, 'Como evitan duplicados en la carga de octubre?', 'Antes de insertar se consulta si existe el identificador funcional. Por ejemplo documento o correo para pacientes, codigo de barras para productos, serial para equipos y una observacion unica para citas.', 'insertarSiNoExiste en seed_agenda_octubre_2026.js')

heading(doc, 'Autenticacion autorizacion y seguridad', 1)
qa(doc, 'Que diferencia hay entre autenticacion y autorizacion?', 'Autenticacion responde quien es el usuario y ocurre en el login. Autorizacion responde que puede hacer; se valida en rutas protegidas usando la especialidad del empleado, especialmente Administracion.', 'empleadoRoutes.js, sesionesAutenticacion.js y autorizacion.js')
qa(doc, 'Como se protegen las contrasenas?', 'El proyecto usa bcryptjs para crear hashes. En el login se compara la contrasena recibida con el hash almacenado. El hash no se expone en las respuestas.', 'contrasenas.js, crearHash y verificarContrasena')
qa(doc, 'Donde se guarda el token?', 'En sessionStorage del navegador. El frontend lo manda como Bearer token en Authorization. Esto permite que las rutas de agenda y cambios administrativos validen la sesion.', 'BIDA_Inicio de sesion_login.js y administracion.js')
qa(doc, 'Por que no basta con ocultar un boton?', 'Ocultar un boton solo mejora la experiencia visual. La seguridad real se valida en el backend: si alguien intenta llamar el endpoint sin sesion o sin el rol correcto, el servidor responde 401 o 403.', 'requiereAdministracion y requiereSesion')
qa(doc, 'Que limitacion tiene la sesion actual?', 'Las sesiones estan en un Map en memoria. Si Render reinicia el proceso, los tokens dejan de existir y el usuario debe iniciar sesion de nuevo. Para una version productiva usaria JWT con rotacion o una sesion persistente en Redis o base de datos.', 'sesionesAutenticacion.js')
qa(doc, 'Que mejora de seguridad es prioritaria?', 'Retirar .env del control de versiones, rotar las credenciales de TiDB y conservar secretos solo en las variables de entorno de Render y en archivos locales ignorados por Git.', 'archivo .env y configuracion de Render')

heading(doc, 'Frontend y experiencia de usuario', 1)
qa(doc, 'Como la pantalla de Administracion obtiene la agenda?', 'administracion.js toma el token de sessionStorage, llama GET /app/cita y renderiza la respuesta. La vista de Administracion puede navegar entre dia, semana y mes usando la lista cargada.', 'cargarAgenda, renderizarAgenda y renderizarCalendarioAgenda')
qa(doc, 'Por que no se debe usar la pagina BIDA_agenda.html para la demostracion?', 'Esa pagina es una maqueta independiente con contenido fijo. La agenda operativa que consulta datos reales esta dentro de administracion.html y se alimenta desde la API.', 'public js BIDA_agenda.html frente a administracion.html')
qa(doc, 'Como se muestran errores del backend al usuario?', 'Cada ruta devuelve un estado HTTP y un mensaje JSON. El frontend revisa response.ok y muestra un mensaje o una fila de error cuando no puede cargar datos.', 'manejo de fetch en administracion.js y respuestas 400 401 403 404 500')
qa(doc, 'Por que se normaliza el body en app.js?', 'El middleware conserva compatibilidad con clientes que puedan enviar JSON como texto o doble codificado. Intenta convertirlo a objeto antes de que llegue a las rutas y responde un 400 claro si no es JSON valido.', 'express.text, express.json y middleware de JSON normalizado')

heading(doc, 'Pruebas y calidad', 1)
qa(doc, 'Que prueban Jest y Supertest?', 'Jest organiza y ejecuta las pruebas. Supertest hace solicitudes HTTP contra la aplicacion Express sin necesitar levantar el servidor en un puerto real.', 'carpeta test y package.json')
qa(doc, 'Por que se usan mocks de mysql2 en las pruebas?', 'Porque una prueba de controlador debe ser repetible y no depender de TiDB, internet ni datos externos. El mock devuelve resultados controlados y permite verificar las respuestas del endpoint.', 'jest.mock mysql2 y mockConnection.query')
qa(doc, 'Que resultado actual puedes citar?', 'La suite actual tiene 7 suites y 48 pruebas aprobadas. Cubre empleados, productos, agenda, ventas, registro, autorizacion y novedades.', 'npm test con Jest')
qa(doc, 'Que no cubren las pruebas actuales?', 'No sustituyen pruebas de navegador ni carga real. Faltaria automatizar el login visual, el calendario, la interaccion completa del usuario y fallos de red simulados con mayor profundidad.')

heading(doc, 'Preguntas de criterio y mejoras', 1)
add_table(doc, ['Pregunta dificil', 'Respuesta tecnica breve', 'Mejora concreta'], [
    ['Que refactor haria primero?', 'Separaria configuracion de base de datos y arranque del servidor en modulos distintos.', 'Crear db.js y server.js para reducir responsabilidad de app.js.'],
    ['Como escalaria el backend?', 'El pool ayuda por instancia, pero las sesiones en memoria no se comparten entre instancias.', 'Usar Redis o JWT y limitar conexiones segun el plan de TiDB.'],
    ['Como evitaria choques de agenda?', 'Hoy se almacenan fecha, hora y profesional, pero debe existir validacion de disponibilidad.', 'Agregar indice unico por idEmpleado fecha y hora y validar antes de crear.'],
    ['Como auditarias cambios?', 'Los CRUD actuales responden resultado, pero no guardan historial detallado.', 'Crear tabla de auditoria con usuario, accion, fecha y datos relevantes.'],
    ['Como protegerias secretos?', 'Las variables de entorno evitan hardcodear credenciales en codigo.', 'Ignorar .env, rotar claves y usar secretos del hosting.']
], [1.55, 2.9, 2.35])

heading(doc, 'Frases utiles para responder con seguridad', 1)
for item in [
    'La regla importante esta en el backend, no solo en la interfaz.',
    'El pool mejora la tolerancia a desconexiones, pero las consultas aun deben manejar errores.',
    'El script de carga es transaccional e idempotente para no dejar datos a medias ni duplicados.',
    'La cobertura actual valida el backend; la siguiente capa de calidad son pruebas de extremo a extremo.',
    'La arquitectura actual es adecuada para el alcance academico y ya identifica mejoras claras para produccion.'
]:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.2); p.paragraph_format.first_line_indent = Inches(-.2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('- ' + item); font(r)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = 'Guia tecnica de preguntas sobre el codigo BIDA'
doc.core_properties.author = 'Proyecto BIDA'
doc.save(OUT)
print(OUT)
