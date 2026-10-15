from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path('output') / 'Guia_sustentacion_BIDA.docx'

BLUE = '17365D'
LIGHT_BLUE = 'DCE6F1'
LIGHT_GRAY = 'F2F2F2'
BORDER = 'D9D9D9'

def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tc_pr.append(shd)

def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in('w:tcBorders')
    if tc_borders is None:
        tc_borders = OxmlElement('w:tcBorders')
        tc_pr.append(tc_borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:{}'.format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn('w:' + side))
        if node is None:
            node = OxmlElement('w:' + side)
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement('w:tblHeader')
    tbl_header.set(qn('w:val'), 'true')
    tr_pr.append(tbl_header)

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('Pagina ')
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText')
    instr_text.set(qn('xml:space'), 'preserve')
    instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)

def style_run(run, bold=False, size=None, color=None):
    run.bold = bold
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def paragraph(doc, text='', style=None, before=0, after=6, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_with_next = keep
    if text:
        r = p.add_run(text)
        style_run(r, size=10.5)
    return p

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    style_run(r, bold=True, size=14 if level == 1 else 11.5, color='000000')
    return p

def qa(doc, question, answer):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run('Pregunta: ' + question)
    style_run(r, bold=True, size=10.5, color=BLUE)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(6)
    p2.paragraph_format.line_spacing = 1.1
    r2 = p2.add_run('Respuesta sugerida: ' + answer)
    style_run(r2, size=10.5)

def table(doc, headers, rows, widths=None):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.style = 'Table Grid'
    header_cells = tbl.rows[0].cells
    for idx, header in enumerate(headers):
        cell = header_cells[idx]
        cell.text = ''
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        style_run(run, bold=True, size=9.5, color='FFFFFF')
        if widths:
            cell.width = Inches(widths[idx])
    set_repeat_table_header(tbl.rows[0])
    for row_index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            cell.text = ''
            if row_index % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(value))
            style_run(run, size=9.2)
            if idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if widths:
                cell.width = Inches(widths[idx])
    return tbl

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.6)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
for name in ['Title', 'Heading 1', 'Heading 2']:
    styles[name].font.name = 'Aptos'
    styles[name]._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    styles[name]._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

footer = sec.footer.paragraphs[0]
footer.paragraph_format.space_before = Pt(4)
add_page_number(footer)
for run in footer.runs:
    style_run(run, size=8, color='666666')

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
run = title.add_run('Guia para la sustentacion de BIDA')
style_run(run, bold=True, size=22, color='000000')

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(18)
run = sub.add_run('Preguntas probables respuestas sugeridas y recorrido de demostracion')
style_run(run, size=11, color='404040')

paragraph(doc, 'Esta guia sirve para presentar BIDA con seguridad ante los profesores. Resume lo que la aplicacion hace hoy, como esta construida y como responder con claridad cuando pregunten por decisiones tecnicas, seguridad, pruebas o limites del proyecto.', after=10)

heading(doc, 'Mensaje de apertura', 1)
paragraph(doc, 'BIDA es una aplicacion web para apoyar la gestion de una clinica odontologica. Centraliza pacientes, empleados, productos, tratamientos, equipos, agenda, ventas y novedades. El objetivo es reducir el manejo disperso de la informacion y permitir que Administracion consulte y gestione los modulos desde un solo lugar.', after=8)

heading(doc, 'Recorrido recomendado para la demostracion', 1)
table(doc, ['Paso', 'Que mostrar', 'Que decir'], [
    ['1', 'Inicio de sesion', 'El acceso identifica al empleado y crea una sesion para consultar los modulos protegidos.'],
    ['2', 'Agenda', 'Como Administracion puedo consultar toda la agenda. Para octubre hay cinco citas por dia y cada cita conserva paciente, profesional, tratamiento y estado.'],
    ['3', 'Pacientes y empleados', 'Los datos se consultan desde la base de datos centralizada. Los empleados tienen especialidad, lo que sirve para asignar las citas.'],
    ['4', 'Productos e inventario', 'Cada producto tiene codigo de barras, precios, categoria, vencimiento y existencias asociadas en inventario.'],
    ['5', 'Tratamientos y equipos', 'Los tratamientos incluyen costo y duracion. Los equipos registran estado, ubicacion y mantenimiento para apoyar la operacion.'],
    ['6', 'Novedades y cierre', 'Las novedades permiten comunicar informacion interna. Cierro mostrando que los datos se conservan en TiDB y la app se despliega en Render.']
], [0.45, 1.55, 4.55])

heading(doc, 'Preguntas sobre el problema y alcance', 1)
qa(doc, 'Que problema resuelve BIDA?', 'Resuelve la necesidad de organizar en una sola aplicacion la informacion operativa de una clinica odontologica. Evita depender de registros separados para pacientes, agenda, inventario, tratamientos y equipos.')
qa(doc, 'Quienes usan el sistema?', 'Principalmente Administracion y el personal de la clinica. Administracion puede gestionar los modulos y visualizar toda la agenda. Los demas empleados consultan las citas que tienen asignadas.')
qa(doc, 'Cuales son los modulos principales?', 'Pacientes, empleados, agenda, productos, inventario, tratamientos, equipos odontologicos, ventas y novedades.')
qa(doc, 'Que funcionalidad considera mas importante?', 'La agenda, porque relaciona paciente, profesional, tratamiento, fecha, hora y estado. Esa informacion permite planear la atencion y dar seguimiento a la operacion diaria.')
qa(doc, 'Por que incluyo datos de demostracion?', 'Para validar el flujo completo y permitir una demostracion realista. Los datos muestran relaciones entre modulos, por ejemplo citas asociadas a pacientes, profesionales y tratamientos.')

heading(doc, 'Preguntas tecnicas', 1)
qa(doc, 'Que tecnologias utiliza el proyecto?', 'El backend esta desarrollado con Node.js y Express. El frontend usa HTML, CSS y JavaScript. La persistencia se maneja con MySQL compatible con TiDB Cloud, y el despliegue se realiza en Render.')
qa(doc, 'Como se comunica el frontend con el backend?', 'El frontend usa fetch para consumir endpoints bajo /app. El backend expone rutas REST para consultar y gestionar cada entidad.')
qa(doc, 'Por que usan TiDB Cloud?', 'Permite tener una base de datos centralizada en la nube, accesible tanto desde el entorno local como desde el despliegue. Su compatibilidad con MySQL facilita el uso de las consultas y del controlador mysql2.')
qa(doc, 'Que solucion implementaron para la perdida de conexion con TiDB?', 'Se sustituyo la conexion unica por un pool de conexiones de mysql2. El pool reutiliza conexiones, mantiene keep alive y puede reemplazar una conexion cerrada temporalmente, reduciendo el riesgo de que el servidor caiga por una desconexion transitoria.')
qa(doc, 'Por que no se modifica la base al iniciar el servidor?', 'La aplicacion se conecta y valida disponibilidad mediante una consulta sencilla. Las cargas de datos se manejan de forma controlada mediante scripts, para evitar cambios inesperados en produccion.')
qa(doc, 'Como manejan las variables de configuracion?', 'El servidor lee host, puerto, usuario, base de datos y SSL desde variables de entorno. En Render se configuran en el panel del servicio y no se deben exponer en el codigo fuente.')

heading(doc, 'Preguntas sobre seguridad y permisos', 1)
qa(doc, 'Como protegen las contrasenas?', 'Las contrasenas nuevas se almacenan con hash bcrypt. Durante el inicio de sesion se verifica el hash, no se compara ni se devuelve la contrasena en texto plano.')
qa(doc, 'Como funciona la autenticacion?', 'Al iniciar sesion el servidor crea un token de sesion. El frontend lo guarda durante la sesion del navegador y lo envía en el encabezado Authorization para acceder a rutas protegidas.')
qa(doc, 'Como controlan los permisos?', 'El backend valida la sesion y consulta la especialidad del empleado. Las acciones administrativas, como crear citas o gestionar ciertos recursos, requieren especialidad Administracion. No se confia solo en ocultar botones en la interfaz.')
qa(doc, 'Que mejoraria en seguridad para una siguiente version?', 'Retiraria cualquier archivo de credenciales del repositorio, rotaria las credenciales expuestas y usaria un gestor de secretos. Tambien migraria las sesiones en memoria a una solucion persistente y agregaria expiracion y renovacion de tokens mas robusta.')

heading(doc, 'Preguntas sobre datos y base de datos', 1)
qa(doc, 'Como se relacionan las tablas?', 'La cita guarda el id del empleado asignado y los datos del paciente. La venta se relaciona con empleado, cliente y detalle de venta. El inventario se relaciona con producto. Las novedades se relacionan con el empleado autor.')
qa(doc, 'Como evitan duplicar los datos de demostracion?', 'Los scripts de carga verifican antes si existe el registro usando identificadores funcionales como documento, correo, codigo de barras, serial u observacion de la cita. Por eso se pueden volver a ejecutar sin crear las mismas citas otra vez.')
qa(doc, 'Cuantos datos hay en la agenda para la demostracion?', 'Se cargaron cinco citas por dia desde el 13 hasta el 31 de octubre de 2026, para un total de 95 citas nuevas. Se distribuyen rotativamente entre los profesionales disponibles.')
qa(doc, 'Que pasa si un empleado no es Administracion?', 'Puede consultar las citas que le corresponden. La consulta se filtra en el servidor por el id del empleado asociado a la sesion.')

heading(doc, 'Preguntas sobre calidad y pruebas', 1)
qa(doc, 'Como probaron el sistema?', 'Se ejecutaron pruebas automatizadas con Jest y Supertest para rutas relevantes de empleados, productos, agenda, ventas, registro, autorizacion y novedades. En la revision actual se aprobaron 48 pruebas en 7 suites.')
qa(doc, 'Que verifican esas pruebas?', 'Verifican respuestas exitosas, validaciones, registros inexistentes, errores de duplicado, autorizacion y operaciones CRUD de los modulos cubiertos.')
qa(doc, 'Que pruebas haria despues?', 'Agregaria pruebas de extremo a extremo desde el navegador, pruebas de carga para el pool de conexiones, validaciones de concurrencia en agenda y una prueba automatizada del despliegue en Render.')
qa(doc, 'Que limitaciones reconoce?', 'La sesion actual vive en memoria del servidor, por lo que un reinicio exige iniciar sesion de nuevo. Tambien se debe fortalecer la gestion de secretos y completar pruebas de interfaz de extremo a extremo.')

heading(doc, 'Respuestas cortas para momentos de presion', 1)
table(doc, ['Si preguntan', 'Respuesta breve'], [
    ['Por que un pool?', 'Porque una sola conexion podia cerrarse y detener el servidor. El pool administra y reutiliza varias conexiones.'],
    ['Los datos son reales?', 'Son datos de demostracion creados para validar los flujos sin usar informacion clinica real.'],
    ['Donde se guardan?', 'En TiDB Cloud, una base de datos centralizada compatible con MySQL.'],
    ['Que ve anar?', 'Es un usuario de Administracion, por lo que puede ver toda la agenda y gestionar los modulos habilitados.'],
    ['Que ocurre si TiDB no responde?', 'La aplicacion informa el error en la consulta; el pool reduce las caidas por conexiones temporales cerradas.'],
    ['Que falta para produccion?', 'Gestion de secretos, sesiones persistentes, pruebas e2e, auditoria y reglas adicionales para agenda e inventario.']
], [1.8, 4.75])

heading(doc, 'Lista final antes de presentar', 1)
checks = [
    'Abrir la URL de Render y confirmar que carga sin errores.',
    'Tener una segunda pestaña con http://localhost:3000 como respaldo.',
    'Iniciar sesion con anar y abrir Administracion luego Agenda.',
    'Ubicar octubre de 2026 y comprobar que se ven cinco citas el dia actual.',
    'Tener lista la pestana de TiDB solo si el profesor solicita evidencia de persistencia.',
    'No mostrar archivos .env ni credenciales durante la presentacion.',
    'Si Render se reinicia, volver a iniciar sesion y continuar desde la agenda.',
    'Cerrar con la mejora futura mas importante: secretos fuera del repositorio y sesiones persistentes.'
]
for item in checks:
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('[ ] ' + item)
    style_run(r, size=10.5)

heading(doc, 'Cierre sugerido', 1)
paragraph(doc, 'BIDA demuestra un flujo funcional de gestion odontologica desde el acceso de usuarios hasta la agenda y los recursos operativos. La aplicacion ya cuenta con persistencia en la nube, control de permisos, datos de demostracion y pruebas automatizadas. Como siguiente etapa, el enfoque seria fortalecer seguridad, sesiones persistentes y pruebas de interfaz para acercarla a un escenario productivo.', after=0)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.core_properties.title = 'Guia para la sustentacion de BIDA'
doc.core_properties.author = 'Proyecto BIDA'
doc.save(OUT)
print(OUT)
